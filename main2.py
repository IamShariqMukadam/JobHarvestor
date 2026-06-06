import pandas as pd
import os
from concurrent.futures import ThreadPoolExecutor, as_completed
from config import CSV_OUTPUT, XLSX_OUTPUT, CHECKPOINT_EVERY
from scrapers.jd_scraper import scrape_jd
from llm.extractor import extract
from utils.driver import get_driver


def load_csv() -> pd.DataFrame:
    df = pd.read_csv(CSV_OUTPUT)
    str_cols = ["Skills Required", "Nice To Have",
                "Experience Required", "Salary", "Status"]
    for col in str_cols:
        if col not in df.columns:
            df[col] = "N/A"
        df[col] = df[col].astype(str).replace("nan", "N/A")
    return df


def already_processed(row) -> bool:
    skills = str(row.get("Skills Required", ""))
    return skills not in ["", "N/A", "nan"]


def process_row(row, naukri_driver=None) -> dict:
    url = str(row.get("URL", ""))
    platform = str(row.get("Platform", ""))
    jd_text = scrape_jd(url, platform, naukri_driver)

    if not jd_text:
        return {
            "Skills Required": "N/A",
            "Nice To Have": "N/A",
            "Experience Required": "N/A",
            "Salary": str(row.get("Salary", "N/A"))
        }

    extracted = extract(jd_text)
    return {
        "Skills Required": ", ".join(extracted["required_skills"]) or "N/A",
        "Nice To Have": ", ".join(extracted["nice_to_have"]) or "N/A",
        "Experience Required": extracted["experience_required"] or "N/A",
        "Salary": extracted["salary"] if extracted["salary"] else str(row.get("Salary", "N/A"))
    }


def store_result(df, idx, result):
    df.at[idx, "Skills Required"] = result["Skills Required"]
    df.at[idx, "Nice To Have"] = result["Nice To Have"]
    df.at[idx, "Experience Required"] = result["Experience Required"]
    df.at[idx, "Salary"] = result["Salary"]


def run():
    print("=" * 55)
    print("  JobHarvestor — Day 2: JD Scraping + LLM Extraction")
    print("=" * 55)

    df = load_csv()
    total = len(df)
    processed = 0
    skipped = 0
    failed = 0

    already_done = df.apply(already_processed, axis=1).sum()
    print(f"\n  Total   : {total}")
    print(f"  Done    : {already_done}")
    print(f"  To do   : {total - already_done}\n")

    # Split by platform
    naukri_rows = [(idx, row) for idx, row in df.iterrows()
                   if row.get("Platform") == "Naukri" and not already_processed(row)]
    other_rows = [(idx, row) for idx, row in df.iterrows()
                  if row.get("Platform") != "Naukri" and not already_processed(row)]

    # --- Internshala + LinkedIn concurrent ---
    def process_other(args):
        idx, row = args
        return idx, row, process_row(row)

    with ThreadPoolExecutor(max_workers=2) as executor:
        futures = {executor.submit(process_other, args): args[0]
                   for args in other_rows}
        for future in as_completed(futures):
            try:
                idx, row, result = future.result()
                store_result(df, idx, result)
                processed += 1
                print(f"  [{idx+1}/{total}] {row.get('Platform')} | "
                      f"{str(row.get('Title',''))[:35]}")
                print(f"    Skills : {result['Skills Required'][:60]}")
                print(f"    Exp    : {result['Experience Required']}")
            except Exception as e:
                print(f"    [!] Failed: {e}")
                failed += 1

            if (processed + failed) % CHECKPOINT_EVERY == 0:
                df.to_csv(CSV_OUTPUT, index=False)
                build_xlsx(df)
                print(f"\n  [Checkpoint] {processed + failed} jobs\n")

    # --- Naukri sequential with shared browser ---
    naukri_driver = get_driver(headless=True)
    print("\n  [Naukri] Shared browser started.")
    try:
        for idx, row in naukri_rows:
            print(f"  [{idx+1}/{total}] Naukri | "
                  f"{str(row.get('Title',''))[:35]}")
            try:
                result = process_row(row, naukri_driver)
                store_result(df, idx, result)
                processed += 1
                print(f"    Skills : {result['Skills Required'][:60]}")
                print(f"    Exp    : {result['Experience Required']}")
            except Exception as e:
                print(f"    [!] Failed: {e}")
                failed += 1

            if (processed + failed) % CHECKPOINT_EVERY == 0:
                df.to_csv(CSV_OUTPUT, index=False)
                build_xlsx(df)
                print(f"\n  [Checkpoint] {processed + failed} jobs\n")
    finally:
        naukri_driver.quit()
        print("  [Naukri] Browser closed.")

    df.drop(columns=["_JD"], errors="ignore").to_csv(CSV_OUTPUT, index=False)
    build_xlsx(df)

    print("\n" + "=" * 55)
    print(f"  Day 2 Complete!")
    print(f"  Processed : {processed}")
    print(f"  Skipped   : {skipped}")
    print(f"  Failed    : {failed}")
    print("=" * 55)


def build_xlsx(df: pd.DataFrame):
    import openpyxl
    from openpyxl.styles import PatternFill, Font, Alignment
    from openpyxl.utils import get_column_letter
    from collections import Counter

    all_skills = []
    for s in df["Skills Required"].dropna():
        if s not in ["N/A", ""]:
            all_skills.extend([x.strip() for x in s.split(",")])

    skill_counts = Counter(all_skills)
    skill_df = pd.DataFrame(skill_counts.most_common(50),
                            columns=["Skill", "Frequency"])
    skill_df["% of Jobs"] = (
        skill_df["Frequency"] / len(df) * 100
    ).round(1).astype(str) + "%"

    platform_df = (df.groupby(["Platform", "Role Searched"])
                   .size().reset_index(name="Job Count"))

    gap_df = pd.DataFrame({
        "Instructions": ["Paste your skills in column B.",
                         "Day 4 will auto-compare against skill frequency."],
        "Your Skills": ["", ""]
    })

    with pd.ExcelWriter(XLSX_OUTPUT, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="All Jobs", index=False)
        platform_df.to_excel(writer, sheet_name="Platform Breakdown", index=False)
        skill_df.to_excel(writer, sheet_name="Skill Frequency", index=False)
        gap_df.to_excel(writer, sheet_name="Gap Analysis", index=False)

    wb = openpyxl.load_workbook(XLSX_OUTPUT)
    colors = {"All Jobs": "1F4E79", "Platform Breakdown": "375623",
              "Skill Frequency": "7B3F00", "Gap Analysis": "4A235A"}
    status_colors = {"To Apply": "D9E1F2", "Applied": "E2EFDA",
                     "Interview": "FFF2CC", "Rejected": "FCE4D6", "Offer": "C6EFCE"}

    for sheet, color in colors.items():
        ws = wb[sheet]
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor=color)
            cell.font = Font(bold=True, color="FFFFFF", size=11)
            cell.alignment = Alignment(horizontal="center")
        for col in ws.columns:
            max_len = max((len(str(c.value)) if c.value else 0 for c in col), default=10)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 60)
        ws.freeze_panes = "A2"

    ws_jobs = wb["All Jobs"]
    cols = list(df.columns)
    if "Status" in cols:
        status_idx = cols.index("Status") + 1
        for row in ws_jobs.iter_rows(min_row=2, max_row=ws_jobs.max_row):
            cell = row[status_idx - 1]
            cell.fill = PatternFill("solid", fgColor=status_colors.get(str(cell.value), "FFFFFF"))

    wb.save(XLSX_OUTPUT)
    print(f"  [XLSX] Updated → {XLSX_OUTPUT}")


if __name__ == "__main__":
    run()