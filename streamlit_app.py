import streamlit as st
import re
import streamlit.components.v1 as stc
import pandas as pd
import plotly.express as px
import json, os, time, sys, subprocess
from collections import Counter
from io import BytesIO
from pathlib import Path

from analysis.gap_analyzer import load_profiles, analyze_gap
from utils.session_store import load_session, save_session
from config import CSV_OUTPUT, XLSX_OUTPUT

import uuid

# ── Per-user session isolation ────────────────────────
if "session_id" not in st.session_state:
    st.session_state.session_id = uuid.uuid4().hex[:10]

SESSION_DIR = f"data/sessions/{st.session_state.session_id}"
os.makedirs(SESSION_DIR, exist_ok=True)
os.environ["JH_SESSION_DIR"] = SESSION_DIR   # inherited by all subprocesses
# ─────────────────────────────────────────────────────

def _cleanup_old_sessions(max_age_hours=24):
    sessions_dir = "data/sessions"
    if not os.path.exists(sessions_dir):
        return
    now = time.time()
    for sid in os.listdir(sessions_dir):
        path = os.path.join(sessions_dir, sid)
        if os.path.isdir(path):
            age = now - os.path.getmtime(path)
            if age > max_age_hours * 3600:
                import shutil
                shutil.rmtree(path, ignore_errors=True)

_cleanup_old_sessions()


st.set_page_config(
    page_title="JobHarvestor",
    page_icon="🎯",
    layout="wide",
    initial_sidebar_state="expanded",
)


if "jh_theme" not in st.session_state:
    st.session_state.jh_theme = "dark"

# if "jh_theme_radio" in st.session_state:
#     _rval = str(st.session_state.jh_theme_radio)
#     st.session_state.jh_theme = "light" if "Light" in _rval else "dark"


_IS_LIGHT = st.session_state.jh_theme == "light"

if _IS_LIGHT:
    _theme_vars = """
      --bg:#F4F3EF; --bg-2:#FFFFFF; --bg-3:#ECEAE3;
      --panel:rgba(255,255,255,.82); --panel-s:rgba(255,255,255,.97);
      --card:rgba(255,255,255,.76); --chat-bg:#FFFFFF;
      --tx:#0C0C0A; --tx-m:rgba(12,12,10,.56); --tx-f:rgba(12,12,10,.32);
      --bd:rgba(12,12,10,.10); --bd-s:rgba(12,12,10,.22);
      --btn-bg:#0C0C0A; --btn-tx:#F4F3EF;
      --inp-bg:#FFFFFF; --inp-hov:#F4F3EF;
      --shadow:0 20px 60px rgba(0,0,0,.10);
      --glow:0 0 0 1px rgba(0,0,0,.07),0 8px 32px rgba(0,0,0,.07);
      --accent:#C9A020; --accent-bg:rgba(201,160,32,.10);
      --grid:rgba(0,0,0,.05);
      --hero-grad:radial-gradient(ellipse at 12% 0%,rgba(201,160,32,.10),transparent 50%),
                 radial-gradient(ellipse at 90% 20%,rgba(0,0,0,.04),transparent 50%);
    """
    _accent_hex = "#E0AA3E"
    _plot_bg = "rgba(255,255,255,.72)"
    _plot_tx = "#050505"
    _plot_grid = "rgba(0,0,0,.12)"
else:
    _theme_vars = """
      --bg:#080808; --bg-2:#040404; --bg-3:#0D0D0D;
      --panel:rgba(12,12,12,.90); --panel-s:rgba(18,18,18,.96);
      --card:rgba(14,14,14,.88); --chat-bg:#060606;
      --tx:#EDE9E0; --tx-m:rgba(237,233,224,.58); --tx-f:rgba(237,233,224,.32);
      --bd:rgba(237,233,224,.09); --bd-s:rgba(237,233,224,.20);
      --btn-bg:#EDE9E0; --btn-tx:#080808;
      --inp-bg:#0D0D0D; --inp-hov:#141414;
      --shadow:0 24px 80px rgba(0,0,0,.82);
      --glow:0 0 0 1px rgba(237,233,224,.08),0 0 40px rgba(237,233,224,.04);
      --accent:#F0C040; --accent-bg:rgba(240,192,64,.10);
      --grid:rgba(255,255,255,.038);
      --hero-grad:radial-gradient(ellipse at 12% 0%,rgba(240,192,64,.12),transparent 50%),
                 radial-gradient(ellipse at 90% 20%,rgba(237,233,224,.05),transparent 50%);
    """
    _accent_hex = "#E0AA3E"
    _plot_bg = "rgba(10,10,10,.84)"
    _plot_tx = "#ffffff"
    _plot_grid = "rgba(255,255,255,.10)"


st.markdown(f"""
<style>
@import url('https://fonts.googleapis.com/css2?family=Syne:wght@400;600;700;800&family=DM+Mono:ital,wght@0,300;0,400;0,500;1,300&family=Inter:wght@300;400;500;600;700&display=swap');

:root {{
{_theme_vars}
  --r:    14px; --r-sm: 9px; --r-lg: 20px; --r-xl: 26px;
  --f-sans: 'Inter', system-ui, sans-serif;
  --f-head: 'Syne', system-ui, sans-serif;
  --f-mono: 'DM Mono', ui-monospace, monospace;
}}
*,*::before,*::after{{box-sizing:border-box}}
html,body,.stApp{{
  background:var(--hero-grad),linear-gradient(180deg,var(--bg) 0%,var(--bg-2) 100%) !important;
  color:var(--tx) !important; font-family:var(--f-sans) !important;
}}
.stApp{{
  --primary-color:var(--tx) !important;
  --background-color:var(--bg) !important;
  --secondary-background-color:var(--inp-bg) !important;
  --text-color:var(--tx) !important;
}}
header[data-testid="stHeader"],div[data-testid="stDecoration"],#MainMenu,footer{{display:none !important}}
.main .block-container{{max-width:1500px !important;padding:0.2rem 2rem 3rem !important}}
h1,h2,h3,h4,h5,h6{{font-family:var(--f-head) !important;color:var(--tx) !important}}
[data-testid="stMarkdownContainer"] p,[data-testid="stMarkdownContainer"] li{{color:var(--tx)}}
.stCaption,[data-testid="stCaptionContainer"]{{color:var(--tx-m) !important}}
hr{{border-color:var(--bd) !important;margin:1.4rem 0 !important}}

/* ── SIDEBAR ── */
section[data-testid="stSidebar"]{{background:var(--bg-2) !important;border-right:1px solid var(--bd) !important}}
section[data-testid="stSidebar"]>div{{background:transparent !important;padding-top:.3rem !important}}
section[data-testid="stSidebar"] [data-testid="stVerticalBlock"]{{gap:.6rem !important}}
/* ── SIDEBAR THEME TOGGLE ── */
section[data-testid="stSidebar"] [data-testid="stToggle"]{{display:flex !important;flex-direction:column !important;align-items:center !important;justify-content:center !important;margin:12px auto 18px !important;width:100% !important}}
section[data-testid="stSidebar"] [data-testid="stToggle"] label{{transform:scale(1.8);transform-origin:center center;display:flex !important;align-items:center !important;justify-content:center !important;gap:10px !important}}
section[data-testid="stSidebar"] [data-testid="stToggle"] p{{font-family:var(--f-mono) !important;font-size:.92rem !important;color:var(--tx-m) !important;text-align:center !important}}


# /* ── SIDEBAR RADIO THEME TOGGLE — hidden, JS-driven ── */
# section[data-testid="stSidebar"] .stRadio{{
#   height:0 !important;overflow:hidden !important;
#   margin:0 !important;padding:0 !important;opacity:0 !important;
# }}
# /* ── SIDEBAR RADIO THEME TOGGLE ── */
# section[data-testid="stSidebar"] .stRadio{{
#   display:flex !important;justify-content:center !important;
# }}
# section[data-testid="stSidebar"] .stRadio>div{{
#   flex-direction:row !important;gap:0 !important;
#   background:var(--bg-3) !important;
#   border:1.5px solid var(--bd-s) !important;
#   border-radius:999px !important;padding:3px !important;
#   margin-bottom:14px !important;display:inline-flex !important;
# }}
# section[data-testid="stSidebar"] .stRadio label{{
#   border-radius:999px !important;padding:6px 20px !important;
#   font-family:var(--f-mono) !important;font-size:.70rem !important;
#   font-weight:600 !important;color:var(--tx-m) !important;
#   cursor:pointer !important;transition:all .2s ease !important;
#   background:transparent !important;border:none !important;
#   display:flex !important;align-items:center !important;
#   gap:5px !important;letter-spacing:.04em !important;
#   user-select:none !important;white-space:nowrap !important;
# }}
# section[data-testid="stSidebar"] .stRadio label > div:first-child{{display:none !important}}
# section[data-testid="stSidebar"] .stRadio label:has(input:checked),
# section[data-testid="stSidebar"] .stRadio label[data-checked="true"]{{
#   background:var(--accent) !important;color:#080808 !important;
#   font-weight:700 !important;box-shadow:0 1px 6px rgba(0,0,0,.25) !important;
# }}
# section[data-testid="stSidebar"] .stRadio input{{display:none !important}}
# section[data-testid="stSidebar"] [data-testid="stRadio"] [data-testid="stMarkdownContainer"]{{display:none !important}}

/* ── SIDEBAR BRAND ── */
.sidebar-brand{{
  position:relative;overflow:hidden;border:1px solid var(--bd-s);border-radius:var(--r-lg);
  padding:17px 17px 15px;margin-bottom:16px;
  background:radial-gradient(ellipse at 0% 0%,var(--accent-bg),transparent 60%),var(--panel-s);
  box-shadow:var(--glow);
}}
.sidebar-brand::before{{
  content:'';position:absolute;inset:0;pointer-events:none;
  background:repeating-linear-gradient(90deg,var(--grid) 0 1px,transparent 1px 32px),
             repeating-linear-gradient(0deg,var(--grid) 0 1px,transparent 1px 32px);
}}
.sb-logo{{position:relative;z-index:1;display:flex;align-items:flex-start;gap:9px;margin-bottom:10px}}
.sb-nameblock{{display:flex;flex-direction:column;gap:3px}}

.sb-icon{{
  width:32px;height:32px;background:var(--accent);color:#080808;border-radius:8px;
  display:flex;align-items:center;justify-content:center;
  font-size:.85rem;font-weight:900;font-family:var(--f-head);flex-shrink:0;
}}
.sb-name{{font-family:var(--f-head);font-size:1rem;font-weight:800;color:var(--tx) !important;line-height:1}}
.sb-sub{{color:var(--tx-m) !important;font-size:.71rem;line-height:1.5}}
.sb-tag{{
  position:relative;z-index:1;display:inline-flex;align-items:center;gap:5px;margin-top:10px;
  border:1px solid var(--accent);color:var(--accent) !important;background:var(--accent-bg);
  border-radius:999px;padding:3px 9px;font-size:.65rem;font-weight:600;
  font-family:var(--f-mono);letter-spacing:.07em;text-transform:uppercase;
}}
.sb-tag::before{{content:'';font-size:.45em}}

/* ── INPUTS ── */
.stTextInput label,.stTextArea label,[data-testid="stSelectbox"] label,[data-testid="stSlider"] label{{
  color:var(--tx) !important;font-weight:600 !important;font-size:.81rem !important;margin-bottom:4px !important;
}}
.stTextInput input,.stTextArea textarea,[data-baseweb="select"]>div{{
  background:var(--inp-bg) !important;color:var(--tx) !important;
  border:1px solid var(--bd-s) !important;border-radius:var(--r-sm) !important;
  box-shadow:none !important;font-family:var(--f-sans) !important;
}}
.stTextInput input:focus,.stTextArea textarea:focus,[data-baseweb="select"]>div:focus-within{{
  border-color:var(--accent) !important;box-shadow:0 0 0 3px var(--accent-bg) !important;
}}
[data-baseweb="select"] div,[data-baseweb="select"] span,[data-baseweb="select"] svg{{
  color:var(--tx) !important;fill:var(--tx) !important;box-shadow:none !important;
}}
[data-baseweb="popover"]>div,[data-baseweb="menu"],div[role="listbox"]{{
  background:var(--panel-s) !important;border:1px solid var(--bd-s) !important;
  border-radius:var(--r) !important;box-shadow:var(--shadow) !important;
}}
[data-baseweb="menu"] li,div[role="option"]{{
  background:transparent !important;color:var(--tx) !important;
  border-radius:7px !important;margin:2px 4px !important;
}}
[data-baseweb="menu"] li:hover,div[role="option"]:hover{{background:var(--inp-hov) !important}}
[data-baseweb="menu"] li[aria-selected="true"],div[role="option"][aria-selected="true"]{{
  background:var(--accent-bg) !important;color:var(--accent) !important;
}}
[data-baseweb="select"]>div,[data-baseweb="select"]>div>div{{
  background:var(--inp-bg) !important;color:var(--tx) !important;
}}

/* ── SLIDER ── */
[data-testid="stSlider"]{{--primary-color:var(--accent) !important}}
[data-testid="stSlider"] [data-baseweb="slider"]{{padding-top:20px !important}}
[data-testid="stSlider"] [role="slider"]{{
  width:20px !important;height:20px !important;
  background:var(--accent) !important;border:2px solid var(--bg) !important;
  border-radius:50% !important;
  box-shadow:0 0 0 2px var(--accent),0 0 14px var(--accent-bg) !important;
}}
/* Track — scoped inside baseweb slider so it never hits label containers */
[data-testid="stSlider"] [data-baseweb="slider"] div[class*="track"],
[data-testid="stSlider"] [data-baseweb="slider"] div[class*="Track"]{{
  background:var(--bd-s) !important;border-radius:999px !important;overflow:hidden !important;
}}
/* Fill portion — explicit accent, overrides --primary-color red fallback */
[data-testid="stSlider"] [data-baseweb="slider"] div[class*="track"]>div:first-child,
[data-testid="stSlider"] [data-baseweb="slider"] div[class*="Track"]>div:first-child{{
  background:var(--accent) !important;border-radius:0 !important;
}}
/* InnerTrack = the actual fill segment — sibling of track, not child */
[data-testid="stSlider"] [data-baseweb="slider"] div[class*="nnerTrack"]{{
  background:var(--accent) !important;border-radius:999px !important;
}}
/* Tick labels — target by testid, not span wildcard bleeding into thumb tooltip */
[data-testid="stSlider"] [data-testid="stTickBarMin"],
[data-testid="stSlider"] [data-testid="stTickBarMax"]{{
  color:var(--tx-m) !important;font-family:var(--f-mono) !important;
  font-size:.72rem !important;background:transparent !important;padding:0 !important;
}}
[data-testid="stSlider"] span{{
  color:var(--tx-m) !important;font-family:var(--f-mono) !important;
  background:transparent !important;padding:0 !important;
}}
:root {{ --primary: var(--accent) !important; }}
[data-testid="stSlider"] div[style*="rgb(255, 75, 75)"] {{
  background: var(--accent) !important;
}}

/* ── BUTTONS ── */
.stButton>button,.stDownloadButton>button,.stLinkButton>a{{
  min-height:40px !important;border-radius:999px !important;
  border:1px solid var(--bd-s) !important;background:var(--panel-s) !important;
  color:var(--tx) !important;font-weight:600 !important;font-size:.81rem !important;
  font-family:var(--f-sans) !important;padding:.42rem 1rem !important;
  box-shadow:none !important;transition:all .16s ease;letter-spacing:.01em;
}}
.stButton>button *,.stDownloadButton>button *,.stLinkButton>a *{{color:inherit !important}}
.stButton>button:hover,.stDownloadButton>button:hover,.stLinkButton>a:hover{{
  background:var(--inp-hov) !important;border-color:var(--accent) !important;
  color:var(--accent) !important;transform:translateY(-1px);
  box-shadow:0 4px 14px var(--accent-bg) !important;
}}
.stButton>button[kind="primary"]{{
  background:var(--btn-bg) !important;color:var(--btn-tx) !important;
  border-color:var(--btn-bg) !important;font-weight:700 !important;box-shadow:var(--glow) !important;
}}
.stButton>button[kind="primary"]:hover{{opacity:.9 !important;transform:translateY(-1px)}}
.stButton>button:disabled{{opacity:.36 !important;transform:none !important;cursor:not-allowed !important}}
[data-testid="stFormSubmitButton"]>button{{
  background:var(--accent) !important;color:#080808 !important;
  border:1px solid var(--accent) !important;font-weight:700 !important;
}}
[data-testid="stFormSubmitButton"]>button:hover{{
  opacity:.85 !important;transform:translateY(-1px) !important;
}}

/* ── HIDE NATIVE SIDEBAR CONTROLS (custom floating btn handles it) ── */
[data-testid="collapsedControl"],[data-testid="stSidebarCollapseButton"],
[data-testid="stSidebarHeader"] button,
section[data-testid="stSidebar"] button[aria-label="Close sidebar"]{{display:none !important}}



/* ── HERO ── */
.jh-hero{{
  position:relative;overflow:hidden;border:1px solid var(--bd-s);border-radius:var(--r-xl);
  padding:38px 40px 34px;margin-bottom:22px;background:var(--panel);
  box-shadow:var(--shadow);backdrop-filter:blur(28px);min-height:228px;
}}
.jh-hero::before{{
  content:'';position:absolute;inset:0;background:var(--hero-grad);pointer-events:none;
}}
.jh-hero::after{{
  content:'';position:absolute;inset:0;pointer-events:none;
  background:repeating-linear-gradient(90deg,var(--grid) 0 1px,transparent 1px 72px),
             repeating-linear-gradient(0deg,var(--grid) 0 1px,transparent 1px 72px);
}}
.jh-hero-inner{{position:relative;z-index:1}}
.jh-kicker{{
  display:inline-flex;align-items:center;gap:8px;
  color:var(--accent) !important;font-family:var(--f-mono);
  font-size:.68rem;font-weight:400;letter-spacing:.13em;text-transform:uppercase;margin-bottom:14px;
}}
.jh-kicker::before{{content:'';width:22px;height:1px;background:var(--accent);display:block}}
.jh-title{{
  font-family:var(--f-head) !important;color:var(--tx) !important;
  font-size:clamp(1.9rem,4.2vw,3.7rem);line-height:.95;font-weight:800;
  max-width:820px;letter-spacing:-.025em;
}}
.jh-title em{{font-style:normal;color:var(--accent) !important}}
.jh-subtitle{{
  color:var(--tx-m) !important;font-size:.92rem;margin-top:13px;
  max-width:660px;line-height:1.65;font-weight:400;
}}
.jh-pills{{display:flex;flex-wrap:wrap;gap:7px;margin-top:20px}}
.jh-pill{{
  border:1px solid var(--bd);background:var(--card);color:var(--tx-m) !important;
  border-radius:999px;padding:4px 11px;font-size:.7rem;font-weight:600;
  font-family:var(--f-mono);backdrop-filter:blur(10px);letter-spacing:.02em;
}}

/* ── TABS ── */
div[data-testid="stTabs"]{{
  background:var(--panel);border:1px solid var(--bd);border-radius:var(--r-lg);
  padding:0 0 1.4rem;box-shadow:var(--shadow);backdrop-filter:blur(28px);overflow:hidden;
}}
div[data-testid="stTabs"]>div:first-child{{
  background:var(--bg-2);border-bottom:1px solid var(--bd);padding:0 1.2rem;margin-bottom:20px;
}}
button[data-baseweb="tab"]{{
  font-family:var(--f-sans) !important;color:var(--tx-m) !important;
  background:transparent !important;border-radius:0 !important;
  border-bottom:2px solid transparent !important;padding:.85rem 1.15rem !important;
  font-weight:600 !important;font-size:.82rem !important;transition:color .15s;
}}
button[data-baseweb="tab"]:hover{{color:var(--tx) !important}}
button[data-baseweb="tab"][aria-selected="true"]{{
  color:var(--accent) !important;border-bottom-color:var(--accent) !important;
}}
[data-baseweb="tab-highlight"],[data-baseweb="tab-border"]{{display:none !important}}
div[data-testid="stTabs"] [data-testid="stVerticalBlock"]{{gap:1rem !important}}

/* ── METRICS ── */
[data-testid="stMetric"]{{
  background:var(--card) !important;border:1px solid var(--bd) !important;
  border-radius:var(--r) !important;backdrop-filter:blur(20px);padding:13px 15px !important;
  position:relative;overflow:hidden;
}}
[data-testid="stMetric"]::before{{
  content:'';position:absolute;top:0;left:0;right:0;height:2px;
  background:linear-gradient(90deg,var(--accent),transparent);opacity:.7;
}}
[data-testid="stMetricLabel"] span{{color:var(--tx-m) !important;font-size:.76rem !important;font-weight:600 !important;letter-spacing:.03em !important}}
[data-testid="stMetricValue"]{{
  color:var(--tx) !important;font-family:var(--f-head) !important;
  font-weight:700 !important;font-size:1.55rem !important;
}}

/* ── EXPANDERS ── */
[data-testid="stExpander"]{{
  background:var(--card) !important;border:1px solid var(--bd) !important;
  border-radius:var(--r) !important;backdrop-filter:blur(20px);
  margin-bottom:9px !important;transition:border-color .18s;
}}
[data-testid="stExpander"]:hover{{border-color:var(--bd-s) !important}}
[data-testid="stExpander"] summary{{
  color:var(--tx) !important;min-height:46px !important;font-weight:500 !important;
  font-family:var(--f-sans) !important;
}}
[data-testid="stExpander"] summary:hover{{color:var(--accent) !important}}
[data-testid="stExpander"] svg{{fill:var(--tx-m) !important}}

/* ── ALERTS ── */
[data-testid="stAlert"]{{border-radius:var(--r-sm) !important;border:1px solid var(--bd) !important}}

/* ── CHAT ── */
[data-testid="stChatMessage"]{{background:transparent !important;padding:.4rem 0 !important}}
[data-testid="stChatInput"]{{
  background:var(--inp-bg) !important;border:1.5px solid var(--bd-s) !important;
  border-radius:var(--r-lg) !important;
  box-shadow:none !important;
  backdrop-filter:blur(24px);padding:6px 6px 6px 10px !important;
  outline:none !important;
}}
[data-testid="stChatInput"]:focus-within{{
  border-color:var(--accent) !important;
  box-shadow:0 0 0 2px var(--accent-bg) !important;
  outline:none !important;
}}
[data-testid="stChatInput"] *:focus{{
  outline:none !important;box-shadow:none !important;border-color:transparent !important;
}}
[data-testid="stChatInput"] [data-baseweb="textarea"],
[data-testid="stChatInput"] [data-baseweb="textarea"]>div,
[data-testid="stChatInput"]>div,[data-testid="stChatInput"]>div>div,
[data-testid="stChatInput"]>div>div>div{{
  background:var(--inp-bg) !important;outline:none !important;border:none !important;
}}
[data-testid="stChatInput"] textarea,[data-testid="stChatInput"] input{{
  color:var(--tx) !important;background:transparent !important;
  caret-color:var(--accent) !important;font-family:var(--f-sans) !important;
  font-size:.92rem !important;line-height:1.5 !important;
  padding:10px 8px !important;
}}
[data-testid="stChatInput"] textarea::placeholder{{color:var(--tx-m) !important}}
[data-testid="stChatInput"] button{{
  background:var(--accent) !important;color:#080808 !important;
  border-radius:12px !important;
  min-width:44px !important;width:44px !important;height:44px !important;
  margin:2px 2px 2px 4px !important;flex-shrink:0 !important;
  display:flex !important;align-items:center !important;justify-content:center !important;
  box-shadow:0 2px 12px var(--accent-bg) !important;
  transition:opacity .15s, transform .15s !important;
}}
[data-testid="stChatInput"] button:hover{{opacity:.85 !important;transform:scale(1.04) !important}}
[data-testid="stChatInput"] button svg{{fill:#080808 !important;width:20px !important;height:20px !important}}

Issue 2 & 7 — Theme toggle broken + Enter flipping mode
BEFORE (lines 24–28):
/* ── PANELS ── */
.agent-panel,.live-scrape-panel,.role-banner{{
  background:var(--card);border:1px solid var(--bd);border-radius:var(--r);
  padding:15px;box-shadow:var(--glow);backdrop-filter:blur(20px);
}}
.agent-panel{{margin-bottom:12px}}
.agent-panel-title,.role-banner-title{{
  font-family:var(--f-head);color:var(--tx) !important;font-weight:700;font-size:.88rem;
}}
.agent-panel-sub{{color:var(--tx-m) !important;font-size:.78rem;margin-top:5px;line-height:1.5}}

.live-scrape-panel{{
  background:var(--card) !important;border:1px solid var(--accent) !important;
  box-shadow:0 0 0 1px var(--accent-bg),0 0 24px var(--accent-bg) !important;
  border-radius:var(--r) !important;padding:14px !important;margin-bottom:14px !important;
}}

/* ── SCRAPE CARD ── */
.scrape-card-outer{{
  margin:16px 0;padding:2px;border-radius:var(--r-lg);
  background:linear-gradient(135deg,var(--bd-s),var(--accent),var(--bd-s),var(--accent));
  background-size:300% 300%;animation:grad-shift 3s ease infinite;
}}
@keyframes grad-shift{{
  0%{{background-position:0% 50%}}50%{{background-position:100% 50%}}100%{{background-position:0% 50%}}
}}
.scrape-card-inner{{
  background:var(--card);border:1px solid var(--bd-s);border-radius:var(--r-lg);
  padding:26px 22px;box-shadow:var(--glow);backdrop-filter:blur(24px);
  position:relative;overflow:hidden;
}}
.scrape-card-inner::before{{
  content:'';position:absolute;inset:0;pointer-events:none;
  background:repeating-linear-gradient(90deg,var(--grid) 0 1px,transparent 1px 48px),
             repeating-linear-gradient(0deg,var(--grid) 0 1px,transparent 1px 48px);
}}
.scrape-title{{
  font-family:var(--f-head);color:var(--tx) !important;
  font-size:1.08rem;font-weight:800;position:relative;z-index:1;margin-bottom:5px;
}}
.scrape-title em{{font-style:normal;color:var(--accent) !important}}
.scrape-sub{{color:var(--tx-m) !important;font-size:.8rem;line-height:1.55;position:relative;z-index:1}}
.scrape-steps{{display:grid;gap:7px;margin-top:16px;position:relative;z-index:1}}
.scrape-step{{
  display:flex;align-items:center;gap:10px;border:1px solid var(--bd);border-radius:var(--r-sm);
  padding:10px 13px;font-size:.81rem;font-weight:500;font-family:var(--f-sans);
  color:var(--tx-f) !important;background:var(--bg-2);transition:all .2s;
}}
.step-num{{
  width:22px;height:22px;border-radius:50%;display:flex;align-items:center;justify-content:center;
  font-size:.66rem;font-weight:700;font-family:var(--f-mono);
  background:var(--bd);color:var(--tx-f);flex-shrink:0;
}}
.scrape-step.active{{
  color:var(--tx) !important;border-color:var(--accent);
  box-shadow:0 0 0 1px var(--accent-bg),0 0 18px var(--accent-bg);
}}
.scrape-step.active .step-num{{background:var(--accent);color:var(--bg)}}
.scrape-step.done{{opacity:.4;text-decoration:line-through;text-decoration-color:var(--tx-f)}}

/* ── PULSE / THINKING ── */
.pulse-row{{display:flex;gap:6px;justify-content:center;margin-top:14px;position:relative;z-index:1}}
.pulse-dot,.thinking-dots span{{
  width:6px;height:6px;border-radius:999px;background:var(--accent);
  display:inline-block;animation:jh-pulse 1.3s infinite ease-in-out;
}}
.thinking-dots span{{margin-right:3px}}
.pulse-dot:nth-child(2),.thinking-dots span:nth-child(2){{animation-delay:.18s}}
.pulse-dot:nth-child(3),.thinking-dots span:nth-child(3){{animation-delay:.36s}}
@keyframes jh-pulse{{
  0%,80%,100%{{opacity:.2;transform:scale(.62)}}
  40%{{opacity:1;transform:scale(1)}}
}}

/* ── LIVE DOT ── */
.live-dot{{
  display:inline-block;width:7px;height:7px;background:#4ADE80;border-radius:50%;
  box-shadow:0 0 0 2px rgba(74,222,128,.22);animation:live-pulse 2s infinite;
}}
@keyframes live-pulse{{
  0%,100%{{box-shadow:0 0 0 2px rgba(74,222,128,.22)}}
  50%{{box-shadow:0 0 0 5px rgba(74,222,128,.07)}}
}}

/* ── TERMINAL ── */
.terminal-box{{
  background:var(--bg-2);border:1px solid var(--bd);border-radius:var(--r-sm);
  padding:13px 15px;max-height:280px;overflow-y:auto;font-family:var(--f-mono) !important;
}}
.terminal-line{{font-family:var(--f-mono);font-size:12.5px;padding:1.5px 0}}
.t-green{{color:#6EE7B7 !important}}.t-blue{{color:#93C5FD !important}}
.t-yellow{{color:var(--accent) !important}}.t-red{{color:#FCA5A5 !important}}
.t-white{{color:var(--tx) !important}}.t-gray{{color:var(--tx-f) !important}}

/* ── BADGES / TAGS ── */
.badge{{
  display:inline-block;padding:3px 9px;border-radius:999px;
  font-size:.68rem;font-weight:600;margin:2px 4px 2px 0;font-family:var(--f-mono);
}}
.badge-status{{background:var(--accent-bg);color:var(--accent) !important;border:1px solid var(--accent)}}
.skill-tag{{
  display:inline-block;background:var(--card);color:var(--tx) !important;
  border-radius:6px;padding:2px 8px;font-size:.7rem;margin:2px;
  border:1px solid var(--bd);font-family:var(--f-mono);
}}

/* ── SECTION HEADER ── */
.section-header{{display:flex;align-items:baseline;gap:10px;margin-bottom:3px}}
.section-title{{font-family:var(--f-head) !important;font-size:1.15rem;font-weight:700;color:var(--tx) !important}}
.section-cap{{font-size:.78rem;color:var(--tx-m) !important;margin-bottom:14px}}

/* ── DATA TABLE ── */
.jh-table-wrap{{
  max-height:340px;overflow:auto;border:1px solid var(--bd);border-radius:var(--r);
  background:var(--card);backdrop-filter:blur(20px);margin-bottom:14px;
}}
.jh-table{{width:100%;border-collapse:collapse;font-size:.83rem}}
.jh-table th{{
  position:sticky;top:0;background:var(--panel-s);color:var(--tx-m) !important;
  padding:10px 14px;text-align:left;border-bottom:1px solid var(--bd-s);
  font-size:.7rem;font-weight:700;letter-spacing:.07em;text-transform:uppercase;font-family:var(--f-mono);
}}
.jh-table td{{padding:10px 14px;border-bottom:1px solid var(--bd);color:var(--tx) !important}}
.jh-table tr:hover td{{background:var(--inp-hov)}}
.rank-num{{font-family:var(--f-mono);font-size:.82rem;color:var(--accent) !important;font-weight:500}}

/* ── PROGRESS ── */
[data-testid="stProgress"]>div>div{{background:var(--accent) !important;border-radius:999px !important}}
[data-testid="stProgress"]>div{{background:var(--bd) !important;border-radius:999px !important;height:3px !important}}

/* ── CHECKBOX ── */
[data-testid="stCheckbox"] label{{color:var(--tx) !important;font-weight:500 !important}}
[data-testid="stCheckbox"] [data-testid="stCheckboxLabel"]{{color:var(--tx) !important}}

/* ── SCORE COLORS ── */
.score-hi{{color:#4ADE80 !important;font-family:var(--f-head);font-size:2rem;font-weight:800}}
.score-mid{{color:var(--accent) !important;font-family:var(--f-head);font-size:2rem;font-weight:800}}
.score-lo{{color:#F87171 !important;font-family:var(--f-head);font-size:2rem;font-weight:800}}

/* ── LIGHT MODE FORCED OVERRIDES ── */
[data-testid="stChatInput"]>div,[data-testid="stChatInput"]>div>div{{
  background:var(--inp-bg) !important;
}}
[data-testid="stBottom"]{{background:var(--bg) !important}}
[data-testid="stBottom"]>div{{background:var(--bg) !important}}

/* Dropdowns / popovers fully themed */
[data-baseweb="popover"],[data-baseweb="popover"]>div,
[data-baseweb="popover"] ul,[data-baseweb="menu"],
div[role="listbox"],div[role="listbox"]>div{{
  background:var(--panel-s) !important;color:var(--tx) !important;
}}
[data-baseweb="popover"] li,div[role="option"]{{
  background:transparent !important;color:var(--tx) !important;
}}
[data-baseweb="popover"] li:hover,div[role="option"]:hover{{
  background:var(--inp-hov) !important;
}}
[data-baseweb="select"]>div,[data-baseweb="select"]>div>div,
[data-baseweb="select"] input{{
  background:var(--inp-bg) !important;color:var(--tx) !important;
}}

/* Expanders — stay themed when open */
[data-testid="stExpander"]>details,
[data-testid="stExpander"]>details[open],
[data-testid="stExpander"]>details>div,
[data-testid="stExpander"] [data-testid="stExpanderDetails"]{{
  background:var(--card) !important;color:var(--tx) !important;
}}
[data-testid="stExpander"]>details>summary,
[data-testid="stExpander"]>details[open]>summary{{
  background:var(--card) !important;color:var(--tx) !important;
}}

/* File uploader */
[data-testid="stFileUploader"],
[data-testid="stFileUploader"]>div,
[data-testid="stFileUploader"]>div>div,
[data-testid="stFileUploadDropzone"],
[data-testid="stFileUploadDropzone"]>div{{
  background:var(--inp-bg) !important;color:var(--tx) !important;
  border-color:var(--bd-s) !important;border-radius:var(--r) !important;
}}
[data-testid="stFileUploaderDropzoneInstructions"] *{{color:var(--tx-m) !important}}
[data-testid="stFileUploadDropzone"] button,
[data-testid="stFileUploader"] button,
[data-testid="stFileUploader"] section button{{
  background:var(--panel-s) !important;border:1px solid var(--bd-s) !important;
  color:var(--tx) !important;border-radius:999px !important;font-weight:600 !important;
}}
section[data-testid="stExpander"] [data-testid="stFileUploader"] *,
section[data-testid="stExpander"] [data-testid="stFileUploadDropzone"] *{{
  background:var(--inp-bg) !important;
}}

@media(max-width:900px){{
  .main .block-container{{padding:1rem !important}}
  .jh-hero{{padding:22px;min-height:190px}}
  .jh-title{{font-size:1.85rem}}
  button[data-baseweb="tab"]{{padding:.75rem .6rem !important;font-size:.76rem !important}}
}}
</style>
""", unsafe_allow_html=True)
# --- INJECT CUSTOM FLOATING SIDEBAR TOGGLE ---
stc.html("""<script>
(function() {
  var MAX_TRIES = 25, tries = 0;

  function init() {
    tries++;
    if (tries > MAX_TRIES) return;

    try {
      var doc = window.parent.document;
      if (!doc || !doc.body) { setTimeout(init, 200); return; }

      /* ── inject button CSS into parent doc once ── */
      if (!doc.getElementById('jh-sb-style')) {
        var s = doc.createElement('style');
        s.id = 'jh-sb-style';
        s.textContent = `
          #jh-sidebar-tog {
            position: fixed;
            left: 0;
            top: 50vh;
            z-index: 99999;
            background: #F0C040;
            color: #080808;
            width: 26px;
            height: 80px;
            display: flex;
            align-items: center;
            justify-content: center;
            cursor: pointer;
            font-size: 20px;
            font-weight: 900;
            border-radius: 0 10px 10px 0;
            user-select: none;
            box-shadow: 3px 0 22px rgba(240,192,64,.35);
            transform: translateY(-50%);
            transition: left .28s cubic-bezier(.4,0,.2,1),
                        background .18s, box-shadow .18s;
          }
          #jh-sidebar-tog:hover {
            background: #E0AA3E;
            box-shadow: 3px 0 32px rgba(240,192,64,.55);
          }
          #jh-sidebar-tog.nudge {
            animation: jh-nudge 1.8s ease-in-out 0.5s 4;
          }
          @keyframes jh-nudge {
            0%,100% { left: var(--jh-base-left, 0px); }
            50%      { left: calc(var(--jh-base-left, 0px) + 7px); }
          }
        `;
        doc.head.appendChild(s);
      }

      /* ── create button div once ── */
      if (doc.getElementById('jh-sidebar-tog')) return;

      var btn = doc.createElement('div');
      btn.id = 'jh-sidebar-tog';

      function getSidebarWidth() {
        var sb = doc.querySelector('section[data-testid="stSidebar"]');
        if (!sb) return 0;
        return parseInt(window.parent.getComputedStyle(sb).width) || 0;
      }

      function updateBtn() {
        var w = getSidebarWidth();
        var isOpen = w > 50;
        btn.textContent = isOpen ? '\u2039' : '\u203a';
        /* slide button to hug the sidebar's right edge */
        btn.style.left = (isOpen ? w : 0) + 'px';
        btn.style.setProperty('--jh-base-left', (isOpen ? w : 0) + 'px');
      }

      /* initial state + nudge hint for closed sidebar */
      updateBtn();
      btn.classList.add('nudge');
      btn.addEventListener('animationend', function() {
        btn.classList.remove('nudge');
      });

      btn.onclick = function() {
        /* briefly un-hide the native button so the click registers with Streamlit */
        var native =
          doc.querySelector('[data-testid="stSidebarCollapseButton"] button') ||
          doc.querySelector('[data-testid="collapsedControl"] button') ||
          doc.querySelector('section[data-testid="stSidebar"] button[aria-label="Close sidebar"]') ||
          doc.querySelector('[data-testid="stSidebarHeader"] button');
        if (native) {
          native.style.cssText = 'display:flex!important';
          native.click();
          setTimeout(function() { native.style.cssText = ''; }, 50);
        }
        /* update position after sidebar animates */
        setTimeout(updateBtn, 150);
        setTimeout(updateBtn, 400);
      };

      doc.body.appendChild(btn);

      /* poll to keep button in sync with sidebar (handles rerun-triggered state changes) */
      setInterval(updateBtn, 500);
      // ── Hide __JH_THEME__ button + expose global click hook ──
      (function(){
        var _themeBtn = null;
        function hideAndWire(){
          try{
            var sb = doc.querySelector('section[data-testid="stSidebar"]');
            if(!sb) return;
            var found = Array.from(sb.querySelectorAll('button')).find(function(b){
              return b.textContent.trim()==='__JH_THEME__';
            });
            if(found){ found.style.cssText='display:none!important'; _themeBtn=found; }
            window.__jhToggleClick = function(){ if(_themeBtn) _themeBtn.click(); };
          }catch(e){}
        }
        hideAndWire();
        var sbEl = doc.querySelector('section[data-testid="stSidebar"]') || doc.body;
        new MutationObserver(hideAndWire).observe(sbEl,{childList:true,subtree:true});
      })();

    } catch(e) { /* cross-origin guard — skip silently */ }
  }

  setTimeout(init, 400);
})();
         

// Fix slider fill color
(function() {
  function fixSliders() {
    try {
      var doc = window.parent.document;
      doc.querySelectorAll('[data-testid="stSlider"] [data-baseweb="slider"] div').forEach(function(d) {
        var c = window.parent.getComputedStyle(d).backgroundColor;
        if (c === 'rgb(255, 75, 75)') {
          d.style.cssText += 'background:#F0C040 !important';
        }
      });
    } catch(e) {}
  }
  fixSliders();
  setInterval(fixSliders, 800);
})();
</script>""", height=0, scrolling=False)


@st.cache_data(ttl=0)
def load_jobs_df(session_dir=None):
    path = os.path.join(session_dir or SESSION_DIR, "JobHarvestor.csv")
    return pd.read_csv(path)


def apply_plotly_theme(fig, height=None):
    kw = dict(
        plot_bgcolor=_plot_bg, paper_bgcolor=_plot_bg,
        font=dict(color=_plot_tx, family="Inter, system-ui, sans-serif", size=12),
        title_font=dict(family="Syne, system-ui, sans-serif", size=14, color=_plot_tx),
        margin=dict(l=24, r=24, t=64, b=42),
        xaxis=dict(
            gridcolor=_plot_grid, zerolinecolor=_plot_grid, color=_plot_tx,
            tickfont=dict(color=_plot_tx), title_font=dict(color=_plot_tx),
        ),
        yaxis=dict(
            gridcolor=_plot_grid, zerolinecolor=_plot_grid, color=_plot_tx,
            tickfont=dict(color=_plot_tx), title_font=dict(color=_plot_tx),
        ),
        legend=dict(bgcolor="rgba(0,0,0,0)", font=dict(color=_plot_tx)),
    )
    if height:
        kw["height"] = height
    fig.update_layout(**kw)
    fig.update_traces(textfont_color=_plot_tx)
    return fig


def persist():
    save_session({
        "skills_input": st.session_state.skills_input,
        "target_tier": st.session_state.target_tier,
        "location": st.session_state.location,
        "custom_roles": st.session_state.custom_roles,
        "jobs_per_platform": st.session_state.jobs_per_plat,
    })


def _scroll_to_bottom():
    pass


def colorize(line):
    l = line.strip()

    if not l:
        return ""

    if "✓" in l or "done" in l.lower() or "complete" in l.lower() or "saved" in l.lower():
        return f'<div class="terminal-line t-green">✔ {l}</div>'

    if "✗" in l or "failed" in l.lower() or "error" in l.lower():
        return f'<div class="terminal-line t-red">✗ {l}</div>'

    if l.startswith("[") and "]" in l:
        tag = l[l.index("[") + 1:l.index("]")]
        rest = l[l.index("]") + 1:]

        color_map = {
            "Internshala": "t-green",
            "LinkedIn": "t-blue",
            "Naukri": "t-yellow",
            "Export": "t-white",
            "CSV": "t-green",
            "XLSX": "t-green",
            "Filter": "t-yellow",
            "Groq": "t-blue",
            "LLM": "t-blue",
            "Embedder": "t-blue",
            "Clusterer": "t-yellow",
            "Analyzer": "t-yellow",
        }

        c = color_map.get(tag, "t-gray")
        return f'<div class="terminal-line"><span class="{c}">[{tag}]</span>{rest}</div>'

    if l.startswith("="):
        return f'<div class="terminal-line t-gray">{l}</div>'

    if "Page" in l or "Offset" in l or "Total" in l:
        return f'<div class="terminal-line t-gray" style="padding-left:16px">{l}</div>'

    return f'<div class="terminal-line t-white">{l}</div>'


def _extract_skills_from_history(chat_msgs):
    user_msgs = [m["content"] for m in chat_msgs if m["role"] == "user"]

    for msg in reversed(user_msgs):
        if "," in msg and len(msg) < 300:
            parts = [p.strip() for p in msg.split(",") if p.strip()]
            if 2 <= len(parts) <= 20:
                return parts

    return []


def _render_gap_chart(gap_data):
    ranked = dict(gap_data.get("all_ranked", []))
    priority = gap_data.get("priority_list", [])
    missing = set(gap_data.get("missing_required", []))

    if not priority:
        return

    rows = [
        {
            "Skill": s,
            "% in JDs": ranked.get(s, 0),
            "Type": "Must Learn" if s in missing else "Useful",
        }
        for s in priority[:12]
    ]

    df_chart = pd.DataFrame(rows)

    fig = px.bar(
        df_chart,
        x="% in JDs",
        y="Skill",
        orientation="h",
        color="Type",
        color_discrete_map={"Must Learn": "#EF5350", "Useful": _accent_hex},
        title=f"Learning Roadmap - Readiness: {gap_data.get('readiness_score', 0)}%",
        text="% in JDs",
    )

    fig.update_traces(texttemplate="%{text}%", textposition="outside")
    fig.update_layout(yaxis={"categoryorder": "total ascending"})
    fig = apply_plotly_theme(fig, height=max(340, len(rows) * 38))

    st.plotly_chart(fig, use_container_width=True)


if "session_loaded" not in st.session_state:
    saved = load_session()
    st.session_state.skills_input = saved.get("skills_input", "")
    st.session_state.target_tier = saved.get("target_tier", "All")
    st.session_state.location = saved.get("location", "India")
    st.session_state.custom_roles = saved.get("custom_roles", [])
    st.session_state.jobs_per_plat = saved.get("jobs_per_platform", 50)
    st.session_state.agent_result = None
    st.session_state.session_loaded = True


with st.sidebar:
    st.markdown("""
<div class="sidebar-brand">
  <div class="sb-logo">
    <div class="sb-icon">JH</div>
    <div class="sb-nameblock">
      <div class="sb-name">JobHarvestor</div>
      <div class="sb-sub">AI Job Market Intelligence</div>
    </div>
  </div>
  <div><span class="sb-tag">Market Intelligence</span></div>
</div>
""", unsafe_allow_html=True)
    # ── Theme toggle ──────────────────────────────────────
    _is_dark = st.session_state.jh_theme == "dark"
    _toggled = st.toggle("🌙  Dark" if _is_dark else "☀️  Light", value=_is_dark, key="jh_theme_tog")
    if _toggled != _is_dark:
        st.session_state.jh_theme = "dark" if _toggled else "light"
        st.rerun()

    st.markdown("**Target Roles**")
    st.markdown('<p style="font-size:.75rem;color:var(--tx);opacity:.75;margin-top:-6px;margin-bottom:10px">Any profession — Data Analyst, Lawyer, Developer...</p>', unsafe_allow_html=True)

    with st.form("add_role_form", clear_on_submit=True):
        new_role = st.text_input("Add a role", placeholder="e.g. Flutter Developer")
        c1, c2 = st.columns(2)
        with c1:
            add_submitted = st.form_submit_button("➕ Add", use_container_width=True)
        with c2:
            clear_submitted = st.form_submit_button("🗑 Clear", use_container_width=True)

    if add_submitted and new_role.strip() and new_role.strip() not in st.session_state.custom_roles:
        st.session_state.custom_roles.append(new_role.strip())
        persist()
        st.rerun()
    if clear_submitted:
        st.session_state.custom_roles = []
        persist()
        st.rerun()

    for i, r in enumerate(st.session_state.custom_roles):
        ca, cb = st.columns([4, 1])

        with ca:
            st.markdown(f"- **{r}**")

        with cb:
            if st.button("x", key=f"del_{i}"):
                st.session_state.custom_roles.pop(i)
                persist()
                st.rerun()

    if not st.session_state.custom_roles:
        st.caption("No roles added yet.")

    st.divider()

    jp = st.slider("Jobs per platform", 10, 100, st.session_state.jobs_per_plat, 10,
                   key="jpp_slider")
    if jp != st.session_state.jobs_per_plat:
        st.session_state.jobs_per_plat = jp
        persist()

    loc = st.text_input(
        "Location",
        value=st.session_state.location,
        placeholder="Mumbai, Bangalore, Remote",
    )

    if loc != st.session_state.location:
        st.session_state.location = loc
        persist()

    tier_opts = ["All", "Mid-market", "Tier-1", "Series A / Startup"]
    tier_idx = tier_opts.index(st.session_state.target_tier) if st.session_state.target_tier in tier_opts else 0

    tt = st.selectbox("Target Company Tier", tier_opts, index=tier_idx)

    if tt != st.session_state.target_tier:
        st.session_state.target_tier = tt
        persist()

    st.markdown("**Your Skills**")

    si = st.text_area(
        "Skills",
        value=st.session_state.skills_input,
        placeholder="Python, SQL, React...",
        height=90,
    )

    if si != st.session_state.skills_input:
        st.session_state.skills_input = si
        persist()

    st.divider()

    keep_old = st.checkbox(
        "Keep previous scrape data",
        value=False,
        help="Uncheck for a fresh scrape each time.",
    )

    st.session_state.keep_old_data = keep_old

    scrape_btn = st.button(
        "⬡  Scrape Jobs + Analyze",
        type="primary",
        use_container_width=True,
        disabled=len(st.session_state.custom_roles) == 0,
    )

    analyze_btn = st.button(
        "▷  Run Agent Analysis",
        use_container_width=True,
        disabled=not os.path.exists(CSV_OUTPUT),
    )

    st.divider()
    st.markdown("**Data Stats**")


    try:
        _s = pd.read_csv(CSV_OUTPUT)
        st.metric("Jobs", len(_s))
        st.metric("Platforms", _s["Platform"].nunique())
        st.metric("Roles", _s["Role Searched"].nunique())
    except Exception:
        st.caption("No data yet.")

    if st.button("↺  Reset Session", use_container_width=True):
        from utils.session_store import SESSION_FILE

        if os.path.exists(SESSION_FILE):
            os.remove(SESSION_FILE)

        load_jobs_df.clear()
        st.rerun()


st.markdown("""
<div class="jh-hero">
  <div class="jh-hero-inner">
    <div class="jh-kicker">Job Market Intelligence Platform</div>
    <div class="jh-title"><em>Job Harvestor</em></div>
    <div class="jh-subtitle">
      Real skill frequencies from scraped JDs across LinkedIn, Naukri &amp; Internshala.
      Conversational AI agent. Personalized gap analysis. Live data — not guesses.
    </div>
    <div class="jh-pills">
      <span class="jh-pill">🕷️ Multi-Platform Scraping</span>
      <span class="jh-pill">🧠 Groq LLM Extraction</span>
      <span class="jh-pill">🗺️ K-Means Clustering</span>
      <span class="jh-pill">📊 Gap Analysis</span>
      <span class="jh-pill">🤖 LangGraph Agent</span>
    </div>
  </div>
</div>
""", unsafe_allow_html=True)


if scrape_btn and st.session_state.custom_roles:
    roles = st.session_state.custom_roles
    n_jobs = st.session_state.jobs_per_plat
    loc_v = st.session_state.location

    st.markdown(f'''<div class="section-header"><span class="section-title">Running Pipeline</span></div><p class="section-cap">{len(roles)} role(s) · 3 platforms · {n_jobs} jobs each</p>''', unsafe_allow_html=True)

    venv_python = sys.executable
    PROJECT_DIR = os.path.dirname(os.path.abspath(__file__))

    os.makedirs("data", exist_ok=True)

    with open(os.path.join(SESSION_DIR, "scrape_config_override.json"), "w") as f:
        json.dump({
            "roles": roles,
            "max_jobs": n_jobs,
            "locations": [l.strip() for l in loc_v.split(",") if l.strip()],
        }, f)

    import config as _c

    if not st.session_state.get("keep_old_data", False):
        for fp in [_c.CSV_OUTPUT, _c.XLSX_OUTPUT]:
            if os.path.exists(fp):
                os.remove(fp)

        load_jobs_df.clear()

    progress = st.progress(0)
    progress_text = st.empty()

    STEPS = [
        ("main.py", "Scraping Platforms", 10, 45),
        ("main2.py", "LLM Skill Extraction", 45, 78),
        ("main3.py", "Embeddings + Clustering", 78, 92),
        ("main4.py", "Building Profiles", 92, 100),
    ]

    pipeline_failed = False

    for cmd, label, p_start, p_end in STEPS:
        progress.progress(p_start)
        progress_text.markdown(f'<span style="font-family:var(--f-mono);font-size:.8rem;color:var(--accent)">{p_start}% ▸ {label}...</span>', unsafe_allow_html=True)

        with st.expander(label, expanded=True):
            out_container = st.empty()
            lines = []

            proc = subprocess.Popen(
                [venv_python, cmd],
                stdout=subprocess.PIPE,
                stderr=subprocess.STDOUT,
                text=True,
                cwd=PROJECT_DIR,
                bufsize=1,
            )

            for line in proc.stdout:
                l = line.rstrip()

                if l:
                    lines.append(colorize(l))
                    html = "".join(lines[-25:])
                    out_container.markdown(
                        f'<div class="terminal-box">{html}</div>',
                        unsafe_allow_html=True,
                    )

            proc.wait()

        progress.progress(p_end)
        progress_text.markdown(f'<span style="font-family:var(--f-mono);font-size:.8rem;color:var(--accent)">{p_end}% ✓ {label} complete</span>', unsafe_allow_html=True)

        if proc.returncode != 0:
            st.error(f"{label} failed. Check the output above.")
            pipeline_failed = True
            break

    if not pipeline_failed:
        load_jobs_df.clear()
        st.cache_data.clear()
        st.session_state.agent_result = None
        st.session_state.post_scrape_roles = roles
        progress_text.markdown('<span style="font-family:var(--f-mono);font-size:.8rem;color:#4ADE80">100% ✓ Pipeline complete</span>', unsafe_allow_html=True)
        st.success("Pipeline complete. Click Run Agent Analysis in the sidebar.")
        time.sleep(3)
        st.rerun()

if analyze_btn:
    venv_python = sys.executable
    PROJECT_DIR = os.path.dirname(os.path.abspath(__file__))

    st.markdown('''<div class="section-header"><span class="section-title">Running Agent Analysis</span></div>''', unsafe_allow_html=True)

    ANALYSIS_STEPS = [
        ("main2.py", "LLM Skill Extraction", 0, 40),
        ("main3.py", "Embeddings + Clustering", 40, 75),
        ("main4.py", "Building Profiles", 75, 100),
    ]

    progress = st.progress(0)
    progress_text = st.empty()
    pipeline_failed = False

    for cmd, label, p_start, p_end in ANALYSIS_STEPS:
        progress.progress(p_start)
        progress_text.markdown(
            f'<span style="font-family:var(--f-mono);font-size:.8rem;color:var(--accent)">{p_start}% ▸ {label}...</span>',
            unsafe_allow_html=True
        )

        with st.expander(label, expanded=True):
            out_container = st.empty()
            lines = []

            proc = subprocess.Popen(
                [venv_python, cmd],
                stdout=subprocess.PIPE,
                stderr=subprocess.STDOUT,
                text=True,
                cwd=PROJECT_DIR,
                bufsize=1,
            )

            for line in proc.stdout:
                l = line.rstrip()
                if l:
                    lines.append(colorize(l))
                    out_container.markdown(
                        f'<div class="terminal-box">{"".join(lines[-25:])}</div>',
                        unsafe_allow_html=True,
                    )

            proc.wait()

        progress.progress(p_end)

        if proc.returncode != 0:
            st.error(f"{label} failed. Check output above.")
            pipeline_failed = True
            break

    if not pipeline_failed:
        load_jobs_df.clear()
        st.cache_data.clear()
        st.session_state.agent_result = None
        progress_text.markdown(
            '<span style="font-family:var(--f-mono);font-size:.8rem;color:#4ADE80">100% ✓ Analysis complete</span>',
            unsafe_allow_html=True
        )
        st.success("Analysis complete! Check Skill Signals and Market Map tabs.")
        time.sleep(2)
        st.rerun()

_BG_RUNNING = os.path.join(SESSION_DIR, ".bg_running")
_BG_DONE    = os.path.join(SESSION_DIR, ".bg_done")
_BG_STOP    = os.path.join(SESSION_DIR, "scraping_stop.flag")
_BG_PHASE   = os.path.join(SESSION_DIR, "scraping_phase.txt")


def _start_background_scrape(role, location="India"):
    if os.path.exists(_BG_RUNNING):
        return False

    os.makedirs("data", exist_ok=True)

    if os.path.exists(_BG_DONE):
        os.remove(_BG_DONE)

    with open(os.path.join(SESSION_DIR, "scrape_config_override.json"), "w") as f:
        json.dump({"roles": [role], "max_jobs": 30, "locations": [location]}, f)

    with open(_BG_RUNNING, "w") as f:
        f.write(role)

    PROJECT_DIR = os.path.dirname(os.path.abspath(__file__))
    _timeouts = [240, 420, 150, 60]

    runner_code = (
        f"import subprocess,sys,os,signal,time\n"
        f"cwd={repr(PROJECT_DIR)}\n"
        f"phases=['scraping','llm','embedding','profiles']\n"
        f"scripts=['main.py','main2.py','main3.py','main4.py']\n"
        f"timeouts={_timeouts}\n"
        f"success=True\n"
        f"for i,s in enumerate(scripts):\n"
        f"    open({repr(_BG_PHASE)},'w').write(phases[i])\n"
        f"    if os.path.exists({repr(_BG_STOP)}): success=False; break\n"
        f"    try:\n"
        f"        proc=subprocess.Popen([sys.executable,s],cwd=cwd,preexec_fn=os.setsid)\n"
        f"        t0=time.time(); killed=False\n"
        f"        while proc.poll() is None:\n"
        f"            time.sleep(2)\n"
        f"            if os.path.exists({repr(_BG_STOP)}) or time.time()-t0>timeouts[i]:\n"
        f"                try: os.killpg(os.getpgid(proc.pid),signal.SIGKILL)\n"
        f"                except Exception: proc.kill()\n"
        f"                proc.wait(); killed=True; break\n"
        f"        if killed or proc.returncode!=0: success=False; break\n"
        f"    except Exception:\n"
        f"        success=False; break\n"
        f"if success:\n"
        f"    open({repr(_BG_DONE)},'w').write('done')\n"
        f"for _f in [{repr(_BG_RUNNING)},{repr(_BG_PHASE)},{repr(_BG_STOP)}]:\n"
        f"    try: os.remove(_f)\n"
        f"    except Exception: pass\n"
    )

    with open("data/bg_runner.py", "w") as f:
        f.write(runner_code)

    subprocess.Popen(
        [sys.executable, "data/bg_runner.py"],
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
        cwd=PROJECT_DIR,
    )

    return True


_KNOWN_ROLES = [
    "data analyst", "data scientist", "data engineer", "business analyst",
    "software engineer", "software developer", "frontend developer",
    "backend developer", "full stack developer", "flutter developer",
    "react developer", "angular developer", "python developer", "java developer",
    "machine learning engineer", "ml engineer", "ai engineer", "product manager",
    "ux designer", "ui designer", "devops engineer", "cloud engineer",
    "qa engineer", "mobile developer", "ios developer", "android developer",
    "content writer", "digital marketer", "project manager", "blockchain developer",
    "lawyer", "legal associate",
]


def _detect_role_quick(text):
    tl = text.lower()

    for r in _KNOWN_ROLES:
        if r in tl:
            return r.title()

    m = re.search(
        r'\b([a-z]+(?:\s+[a-z]+)?)\s+'
        r'(developer|analyst|engineer|manager|designer|scientist|architect)\b',
        tl,
    )

    if m and m.group(1) not in ("a", "the", "any", "software", "i"):
        return (m.group(1) + " " + m.group(2)).title()

    m2 = re.search(r'(?:analyze|search|about|find|for)\s+([a-z]+(?:\s+[a-z]+)?)\s+jobs?', tl)

    if m2:
        r = m2.group(1).strip()

        if 1 <= len(r.split()) <= 3:
            return r.title()

    return None


_KNOWN_LOCATIONS = [
    "pune", "mumbai", "bangalore", "bengaluru", "delhi", "hyderabad", "chennai",
    "kolkata", "noida", "gurgaon", "gurugram", "ahmedabad", "india", "remote",
    "work from home", "wfh", "pan india", "new delhi", "navi mumbai", "thane",
    "kochi", "jaipur", "bhopal", "indore", "surat", "vadodara", "lucknow",
]


def _detect_location_quick(text):
    tl = text.lower()

    for loc in _KNOWN_LOCATIONS:
        if loc in tl:
            if loc in ("remote", "work from home", "wfh"):
                return "Remote"

            if loc in ("india", "pan india"):
                return "India"

            return loc.title()

    m = re.search(r'\b(?:in|at|for|near|around)\s+([a-z]+(?:\s+[a-z]+)?)', tl)

    if m:
        cand = m.group(1).strip()

        if 1 <= len(cand.split()) <= 2 and cand not in ("the", "a", "an", "my", "your", "india"):
            return cand.title()

    return None


tab1, tab2, tab3, tab4 = st.tabs([
    "Agent Console",
    "Skill Signals",
    "Market Map",
    "Job Pipeline",
])


with tab1:
    from agent.graph import run_chat_turn, GREETING

    if "chat_msgs" not in st.session_state:
        st.session_state.chat_msgs = []
        st.session_state.agent_history = []

    if os.path.exists(_BG_DONE):
        os.remove(_BG_DONE)
        load_jobs_df.clear()
        st.cache_data.clear()

        _role = st.session_state.pop("bg_scrape_role", "the role")

        st.session_state.chat_msgs.append({
            "role": "assistant",
            "content": (
                f"Fresh data is ready. I scraped and processed **{_role}** job listings.\n\n"
                f"Shall I run the full skill frequency analysis and gap report?"
            ),
        })

    scraping_active = os.path.exists(_BG_RUNNING)
    col_chat, col_info = st.columns([3, 1], gap="large")

    with col_info:
        if scraping_active:
            with open(_BG_RUNNING) as _f:
                _rr = _f.read().strip()

            st.markdown(f"""
<div class="live-scrape-panel">
  <div style="display:flex;align-items:center;gap:8px;margin-bottom:6px">
    <span class="live-dot"></span>
    <span style="font-family:var(--f-head);font-weight:700;font-size:.87rem">Live Market Scan</span>
  </div>
  <div style="font-size:.78rem;color:var(--tx-m);font-family:var(--f-mono)">{_rr}</div>
  <div class="pulse-row"><span class="pulse-dot"></span><span class="pulse-dot"></span><span class="pulse-dot"></span></div>
</div>
""", unsafe_allow_html=True)

        st.markdown("""
<div class="agent-panel">
  <div class="agent-panel-title">Quick Prompts</div>
  <div class="agent-panel-sub">Tap to send a market intelligence query instantly.</div>
</div>
""", unsafe_allow_html=True)

        suggestions = [
            "Analyze Data Analyst jobs",
            "What skills does FAANG look for?",
            "I know Python, SQL, Tableau",
            "What's trending in ML roles?",
            "Compare startup vs mid-market",
        ]

        for sug in suggestions:
            if st.button(sug, use_container_width=True, key=f"sug_{hash(sug)}", disabled=scraping_active):
                st.session_state._pending_input = sug
                st.rerun()

        st.divider()

        if st.button("↺  New Conversation", use_container_width=True):
            st.session_state.chat_msgs = []
            st.session_state.agent_history = []

            for _flag in [_BG_RUNNING, _BG_DONE, _BG_STOP, _BG_PHASE]:
                try:
                    os.remove(_flag)
                except Exception:
                    pass

            for _key in ["bg_scrape_role", "chat_pending_role", "chat_detected_location"]:
                st.session_state.pop(_key, None)

            st.rerun()

    with col_chat:
        st.markdown('''<div class="section-header"><span class="section-title">JobHarvestor Agent</span></div><p class="section-cap">Conversational market analysis · live scraping · skill extraction · gap reporting</p>''', unsafe_allow_html=True)

        if "_pending_role" in st.session_state:
            pr = st.session_state["_pending_role"]
            existing = [r for r in st.session_state.custom_roles if r.lower() != pr.lower()]

            st.markdown(f"""
<div class="role-banner">
  <div class="role-banner-title">New role detected — {pr}</div>
  <div style="color:var(--tx-m);font-size:.8rem;margin-top:5px">
    You have <b>{len(existing)}</b> existing role(s). Choose how to continue.
  </div>
</div>
""", unsafe_allow_html=True)

            bc1, bc2, bc3 = st.columns([5, 5, 2], gap="medium")

            with bc1:
                if st.button(f"Keep existing + add {pr}", use_container_width=True):
                    if pr not in st.session_state.custom_roles:
                        st.session_state.custom_roles.append(pr)
                        persist()

                    st.session_state["_pending_input"] = st.session_state.pop("_pending_msg", "")
                    del st.session_state["_pending_role"]
                    st.rerun()

            with bc2:
                if st.button(f"Clear all and focus on {pr}", use_container_width=True):
                    st.session_state.custom_roles = [pr]
                    persist()
                    st.session_state.agent_history = []
                    st.session_state.chat_msgs = []
                    st.session_state["_pending_input"] = st.session_state.pop("_pending_msg", "")
                    del st.session_state["_pending_role"]
                    st.rerun()

            with bc3:
                if st.button("Cancel", use_container_width=True):
                    st.session_state.pop("_pending_role", None)
                    st.session_state.pop("_pending_msg", None)
                    st.rerun()

            st.stop()

        if not st.session_state.chat_msgs:
            st.session_state.chat_msgs.append({"role": "assistant", "content": GREETING})

        if scraping_active:
            with open(_BG_RUNNING) as _f:
                _running_role = _f.read().strip()

            _phase = "scraping"

            if os.path.exists(_BG_PHASE):
                with open(_BG_PHASE) as _pf:
                    _phase = _pf.read().strip()

            _phase_idx = {"scraping": 0, "llm": 1, "embedding": 2, "profiles": 3}
            _cur = _phase_idx.get(_phase, 0)

            _steps_data = [
                ("01", "Scanning LinkedIn, Naukri & Internshala"),
                ("02", "LLM skill signal extraction via Groq"),
                ("03", "Embedding + clustering market data"),
                ("04", "Building tier demand profiles"),
            ]

            _steps_html = ""

            for _i, (_num, _lbl) in enumerate(_steps_data):
                _cls = "done" if _i < _cur else ("active" if _i == _cur else "")
                _steps_html += f'<div class="scrape-step {_cls}"><span class="step-num">{_num}</span>{_lbl}</div>'

            for msg in st.session_state.chat_msgs:
                with st.chat_message(msg["role"]):
                    st.markdown(msg["content"])

            st.markdown(f"""
<div class="scrape-card-outer">
  <div class="scrape-card-inner">
    <div class="scrape-title">Processing <em>{_running_role}</em> market data</div>
    <div class="scrape-sub">Background pipeline running. Chat resumes automatically when complete.</div>
    <div class="scrape-steps">{_steps_html}</div>
    <div class="pulse-row">
      <span class="pulse-dot"></span><span class="pulse-dot"></span><span class="pulse-dot"></span>
    </div>
  </div>
</div>
""", unsafe_allow_html=True)

            _sc1, _sc2, _sc3 = st.columns([2, 1, 2])
            with _sc2:
                if st.button("⏹ Stop", key="stop_real_btn", type="secondary", use_container_width=True):
                    Path(_BG_STOP).touch()
                    for _fl in [_BG_RUNNING, _BG_PHASE, _BG_DONE]:
                        if os.path.exists(_fl): os.remove(_fl)
                    for _key in ["bg_scrape_role", "chat_pending_role", "chat_detected_location"]:
                        st.session_state.pop(_key, None)
                    st.rerun()

            st.chat_input("Market scan is running. Please wait...", disabled=True)
            time.sleep(5)
            st.rerun()

            st.markdown(f"""
<div class="scrape-card-outer">
  <div class="scrape-card-inner">
    <div class="scrape-title">Processing <em>{_running_role}</em> market data</div>
    <div class="scrape-sub">Background pipeline running. Chat resumes automatically when complete.</div>
    <div class="scrape-steps">{_steps_html}</div>
    <div class="pulse-row">
      <span class="pulse-dot"></span><span class="pulse-dot"></span><span class="pulse-dot"></span>
    </div>
  </div>
</div>
""", unsafe_allow_html=True)

            st.chat_input("Market scan is running. Please wait...", disabled=True)
            time.sleep(5)
            st.rerun()

        else:
            for msg in st.session_state.chat_msgs:
                with st.chat_message(msg["role"]):
                    st.markdown(msg["content"])

            _scroll_to_bottom()

            pending = st.session_state.pop("_pending_input", None)
            user_input = st.chat_input("Ask about any role, your skills, or the market…") or pending

            if user_input:
                with st.chat_message("user"):
                    st.markdown(user_input)

                st.session_state.chat_msgs.append({"role": "user", "content": user_input})

                _detected = _detect_role_quick(user_input)
                _detected_loc = _detect_location_quick(user_input)

                if _detected_loc:
                    st.session_state.chat_detected_location = _detected_loc

                if _detected and not st.session_state.get("chat_detected_location"):
                    st.session_state.chat_pending_role = _detected

                if (
                    not _detected
                    and st.session_state.get("chat_pending_role")
                    and st.session_state.get("chat_detected_location")
                ):
                    _detected = st.session_state.chat_pending_role

                _existing = st.session_state.custom_roles
                _already = _detected and any(r.lower() == _detected.lower() for r in _existing)

                if _detected and _existing and not _already and not os.path.exists(_BG_RUNNING):
                    st.session_state.chat_msgs.pop()
                    st.session_state["_pending_role"] = _detected
                    st.session_state["_pending_msg"] = user_input
                    st.rerun()

                _final_loc = st.session_state.get("chat_detected_location", "")

                if not os.path.exists(_BG_RUNNING) and _detected and _final_loc:
                    _started = _start_background_scrape(_detected, _final_loc)

                    if _started:
                        st.session_state["bg_scrape_role"] = _detected

                        if _detected not in st.session_state.custom_roles:
                            st.session_state.custom_roles.append(_detected)
                            persist()

                        st.session_state.pop("chat_pending_role", None)
                        st.session_state.pop("chat_detected_location", None)

                        _ack = (
                            f"Got it. I'm starting a live market scan for **{_detected}** jobs in "
                            f"**{_final_loc}**.\n\n"
                            f"I'll fetch listings, extract skill signals, cluster the market data, and resume chat "
                            f"when the intelligence layer is ready."
                        )

                        st.session_state.chat_msgs.append({"role": "assistant", "content": _ack})
                        st.session_state.agent_history.append({"role": "user", "content": user_input})
                        st.session_state.agent_history.append({"role": "assistant", "content": _ack})
                        st.rerun()

                with st.chat_message("assistant"):
                    thinking = st.empty()

                    thinking.markdown(
                        '<div class="thinking-dots"><span></span><span></span><span></span></div>',
                        unsafe_allow_html=True,
                    )

                    try:
                        response_text, updated_history = run_chat_turn(
                            user_input,
                            st.session_state.agent_history,
                        )

                        st.session_state.agent_history = updated_history

                        gap_data = None

                        if any(
                            kw in response_text.lower()
                            for kw in ["readiness score", "missing skills", "learning roadmap", "% of jds"]
                        ):
                            try:
                                _mentioned = _extract_skills_from_history(st.session_state.chat_msgs)

                                if _mentioned:
                                    gap_data = analyze_gap(
                                        _mentioned,
                                        st.session_state.get("target_tier", "Mid-market"),
                                    )
                            except Exception:
                                pass

                        thinking.markdown(response_text)

                        if gap_data and gap_data.get("priority_list"):
                            _render_gap_chart(gap_data)

                    except Exception as e:
                        err = f"**Agent error:** {e}\n\nCheck `GROQ_API_KEY` in `.env`."
                        thinking.markdown(err)
                        response_text = err

                st.session_state.chat_msgs.append({"role": "assistant", "content": response_text})
                _scroll_to_bottom()
                st.rerun()


with tab2:
    st.markdown('''<div class="section-header"><span class="section-title">Skill Signals</span></div><p class="section-cap">What skills actually appear in JDs — real demand vs boilerplate filler</p>''', unsafe_allow_html=True)

    try:
        df = load_jobs_df(SESSION_DIR)

        c1, c2, c3 = st.columns(3, gap="large")

        with c1:
            ra = ["All"] + sorted(df["Role Searched"].dropna().unique().tolist())
            post = st.session_state.get("post_scrape_roles", [])
            def_r = post[0] if post and post[0] in ra else "All"
            sr = st.selectbox("Filter by Role", ra, index=ra.index(def_r), key="hm_role")

        with c2:
            pa = ["All"] + sorted(df["Platform"].dropna().unique().tolist())
            sp = st.selectbox("Filter by Platform", pa, key="hm_plat")

        with c3:
            tn = st.slider("Top N skills", 10, 50, 25, key="hm_n")

        st.markdown("")

        df_f = df.copy()

        if sr != "All":
            df_f = df_f[df_f["Role Searched"] == sr]

        if sp != "All":
            df_f = df_f[df_f["Platform"] == sp]

        skills_all = []

        for s in df_f["Skills Required"].dropna():
            if str(s) not in ["N/A", "nan", ""]:
                skills_all.extend([x.strip() for x in str(s).split(",") if x.strip()])

        if skills_all:
            counts = Counter(skills_all)
            sk_df = pd.DataFrame(counts.most_common(tn), columns=["Skill", "Count"])
            sk_df["% of JDs"] = (sk_df["Count"] / len(df_f) * 100).round(1)

            color_scale = (
                [[0, "rgba(201,160,32,.15)"], [1, _accent_hex]]
                if _IS_LIGHT
                else [[0, "rgba(240,192,64,.15)"], [1, _accent_hex]]
            )

            fig = px.bar(
                sk_df,
                x="% of JDs",
                y="Skill",
                orientation="h",
                color="% of JDs",
                color_continuous_scale=color_scale,
                title=f"Top {tn} Skill Signals - {sr} ({len(df_f)} jobs)",
                text="% of JDs",
            )

            fig.update_traces(texttemplate="%{text}%", textposition="outside")
            fig.update_layout(yaxis={"categoryorder": "total ascending"}, showlegend=False)
            fig = apply_plotly_theme(fig, height=max(480, tn * 26))
            fig.update_layout(coloraxis_showscale=False)

            st.plotly_chart(fig, use_container_width=True)

            st.markdown("#### Signal Frequency Table")

            rows_html = "".join(
                f'<tr>'
                f'<td><span class="rank-num">#{i+1:02d}</span></td>'
                f'<td style="font-weight:600">{r["Skill"]}</td>'
                f'<td style="font-family:var(--f-mono);text-align:right">{r["Count"]}</td>'
                f'<td style="text-align:right"><span style="font-family:var(--f-mono);color:var(--accent);font-weight:500">{r["% of JDs"]}%</span></td>'
                f'</tr>'
                for i, (_, r) in enumerate(sk_df.iterrows())
            )

            st.markdown(f"""
<div class="jh-table-wrap">
  <table class="jh-table">
    <thead>
      <tr>
        <th>#</th>
        <th>Skill</th>
        <th>Count</th>
        <th>% of JDs</th>
      </tr>
    </thead>
    <tbody>{rows_html}</tbody>
  </table>
</div>
""", unsafe_allow_html=True)

            c_dl1, c_dl2 = st.columns(2, gap="large")

            with c_dl1:
                st.download_button(
                    "Download CSV",
                    data=sk_df.to_csv(index=False),
                    file_name="skill_frequency.csv",
                    mime="text/csv",
                    use_container_width=True,
                )

            with c_dl2:
                buf2 = BytesIO()

                with pd.ExcelWriter(buf2, engine="openpyxl") as w:
                    sk_df.to_excel(w, sheet_name="Skill Frequency", index=False)
                    df_f.to_excel(w, sheet_name="All Jobs", index=False)

                buf2.seek(0)

                st.download_button(
                    "Download XLSX",
                    data=buf2.read(),
                    file_name="skill_frequency.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True,
                )

        else:
            st.warning("No skill data. Run the extraction pipeline first.")

    except FileNotFoundError:
        st.warning("No data yet. Add roles and scrape.")

    except Exception as e:
        st.error(f"Error: {e}")


with tab3:
    st.markdown('''<div class="section-header"><span class="section-title">Market Map</span></div><p class="section-cap">Each dot = one job. Similar skill requirements cluster together.</p>''', unsafe_allow_html=True)
    st.caption("Each dot is one job. Similar skill requirements appear closer together.")

    try:
        df = load_jobs_df(SESSION_DIR)

        if "PCA_X" not in df.columns:
            st.warning("Cluster data not found. Run Scrape Jobs + Analyze to generate it.")

        else:
            df["Cluster"] = pd.to_numeric(df["Cluster"], errors="coerce").fillna(0).astype(int)

            cluster_labels = {}

            try:
                for p in load_profiles():
                    cluster_labels[p["cluster_id"]] = p["label"]
            except Exception:
                pass

            df["Cluster Label"] = df["Cluster"].map(cluster_labels).fillna("Mid-market")

            c1, c2, c3 = st.columns(3, gap="large")

            with c1:
                co = ["All"] + sorted(df["Cluster Label"].unique().tolist())
                sc = st.selectbox("Filter by Tier", co)

            with c2:
                ro = ["All"] + sorted(df["Role Searched"].dropna().unique().tolist())
                src = st.selectbox("Filter by Role", ro, key="cm_role")

            with c3:
                ac = sorted(df["Company"].dropna().unique().tolist())
                sel_co = st.multiselect("Highlight Companies", ac, placeholder="Select companies...")

            df_c = df.copy()

            if sc != "All":
                df_c = df_c[df_c["Cluster Label"] == sc]

            if src != "All":
                df_c = df_c[df_c["Role Searched"] == src]

            df_c["Highlight"] = df_c["Company"].apply(lambda x: x if x in sel_co else "Other")
            cc = "Highlight" if sel_co else "Cluster Label"

            fig = px.scatter(
                df_c,
                x="PCA_X",
                y="PCA_Y",
                color=cc,
                hover_data={
                    "Title": True,
                    "Company": True,
                    "Platform": True,
                    "Location": True,
                    "PCA_X": False,
                    "PCA_Y": False,
                },
                color_discrete_map={
                    "Tier-1": _accent_hex,
                    "Series A / Startup": "#FA410E",
                    "Mid-market": "#70eb18",
                    "Other": "rgba(150,150,150,.35)",
                },
                title=f"Market Clusters ({len(df_c)} jobs)",
                labels={
                    "PCA_X": "Skill Cluster Axis 1",
                    "PCA_Y": "Skill Cluster Axis 2",
                },
            )

            fig.update_traces(marker=dict(size=7, opacity=0.78))
            fig = apply_plotly_theme(fig, height=540)

            st.plotly_chart(fig, use_container_width=True)

            for label, group in df_c.groupby("Cluster Label"):
                with st.expander(f"{label} - {len(group)} jobs"):
                    ca2, cb2 = st.columns(2, gap="large")

                    with ca2:
                        st.markdown("**Top Companies**")

                        for co_name, cnt in group["Company"].value_counts().head(8).items():
                            st.markdown(f"- {co_name} ({cnt})")

                    with cb2:
                        st.markdown("**Top Skills**")

                        cs = []
                        _bad = {"n/a", "nan", "null", "none", ""}

                        for s in group["Skills Required"].dropna():
                            if str(s).strip().lower() not in _bad:
                                cs.extend([
                                    x.strip()
                                    for x in str(s).split(",")
                                    if x.strip().lower() not in _bad
                                ])

                        for sk, cnt in Counter(cs).most_common(8):
                            pct = round(cnt / len(group) * 100, 1)
                            st.markdown(
                                f'• **{sk}** — <span style="color:var(--accent);font-family:var(--f-mono);font-weight:500">{pct}%</span>',
                                unsafe_allow_html=True
                            )

    except FileNotFoundError:
        st.warning("No data yet.")

    except Exception as e:
        st.error(f"Error: {e}")


with tab4:
    st.markdown('''<div class="section-header"><span class="section-title">Job Pipeline</span></div><p class="section-cap">Track applications, update status, and export your pipeline</p>''', unsafe_allow_html=True)

    PLATFORM_COLORS = {
        "LinkedIn": "#777777",
        "Naukri": "#777777",
        "Internshala": "#777777",
    }

    STATUS_OPTIONS = ["To Apply", "Applied", "Interview", "Rejected", "Offer", "Automate"]

    STATUS_EMOJI = {
        "To Apply": "○",
        "Applied": "●",
        "Interview": "◐",
        "Rejected": "×",
        "Offer": "★",
        "Automate": "◆",
    }

    try:
        df = load_jobs_df(SESSION_DIR)

        fc1, fc2, fc3 = st.columns(3, gap="large")

        with fc1:
            rlist = ["All"] + sorted(df["Role Searched"].dropna().unique().tolist())
            post2 = st.session_state.get("post_scrape_roles", [])
            defr2 = post2[0] if post2 and post2[0] in rlist else "All"
            rf = st.selectbox("Role", rlist, index=rlist.index(defr2), key="jb_role")

        with fc2:
            plist = ["All"] + sorted(df["Platform"].dropna().unique().tolist())
            pf = st.selectbox("Platform", plist, key="jb_plat")

        with fc3:
            sf = st.selectbox("Status", ["All"] + STATUS_OPTIONS, key="jb_status")

        df_j = df.copy()

        if rf != "All":
            df_j = df_j[df_j["Role Searched"] == rf]

        if pf != "All":
            df_j = df_j[df_j["Platform"] == pf]

        if sf != "All":
            df_j = df_j[df_j["Status"] == sf]

        sc1, sc2, sc3, sc4, sc5, sc6 = st.columns([1, 1, 1, 1, 1.45, 1.45], gap="large")

        with sc1:
            st.metric("Total", len(df_j))

        with sc2:
            st.metric("To Apply", len(df_j[df_j["Status"] == "To Apply"]))

        with sc3:
            st.metric("Applied", len(df_j[df_j["Status"] == "Applied"]))

        with sc4:
            st.metric("Interview", len(df_j[df_j["Status"] == "Interview"]))

        with sc5:
            buf3 = BytesIO()

            with pd.ExcelWriter(buf3, engine="openpyxl") as w:
                df_j.to_excel(w, sheet_name="Filtered Jobs", index=False)
                df.to_excel(w, sheet_name="All Jobs", index=False)

            buf3.seek(0)

            st.download_button(
                "Export XLSX",
                data=buf3.read(),
                file_name="jobs_export.xlsx",
                use_container_width=True,
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )

        with sc6:
            with st.expander("Upload XLSX / CSV"):
                up = st.file_uploader("Upload", type=["xlsx", "csv"], label_visibility="collapsed")

                if up:
                    try:
                        df_up = pd.read_csv(up) if up.name.endswith(".csv") else pd.read_excel(up, sheet_name="All Jobs")
                        df_up.to_csv(CSV_OUTPUT, index=False)
                        load_jobs_df.clear()
                        st.success("Loaded.")
                        st.rerun()

                    except Exception as e:
                        st.error(str(e))

        st.divider()

        changed = False

        for loop_i, (_, row) in enumerate(df_j.head(100).iterrows()):
            ridx = df[df["URL"] == row["URL"]].index.tolist()
            ridx = ridx[0] if ridx else None

            plat = str(row.get("Platform", ""))
            pc = PLATFORM_COLORS.get(plat, "#777777")
            stat = str(row.get("Status", "To Apply"))
            emoji = STATUS_EMOJI.get(stat, "○")

            with st.expander(
                f"{emoji} {row.get('Title', 'N/A')} - "
                f"{row.get('Company', 'N/A')} - "
                f"{row.get('Location', 'N/A')}"
            ):
                d1, d2 = st.columns([3, 1], gap="large")

                with d1:
                    st.markdown(
                        f'<span class="badge" style="background:{pc};color:white">{plat}</span>'
                        f'<span class="badge badge-status">{row.get("Date Posted", "N/A")}</span>',
                        unsafe_allow_html=True,
                    )

                    sk = str(row.get("Skills Required", "N/A"))

                    if sk not in ["N/A", "nan", ""]:
                        tags = "".join([
                            f'<span class="skill-tag">{s.strip()}</span>'
                            for s in sk.split(",")[:8]
                            if s.strip()
                        ])

                        st.markdown(f"**Skills:** {tags}", unsafe_allow_html=True)

                    e_val = row.get("Experience Required", "N/A")
                    s_val = row.get("Salary", "N/A")

                    st.markdown(
                        f"**Exp:** {e_val} &nbsp;|&nbsp; **Salary:** {s_val}",
                        unsafe_allow_html=True,
                    )

                    url = str(row.get("URL", ""))

                    if url not in ["N/A", "nan", ""]:
                        st.link_button("Apply Now", url)

                with d2:
                    cur_idx = STATUS_OPTIONS.index(stat) if stat in STATUS_OPTIONS else 0

                    new_s = st.selectbox(
                        "Status",
                        STATUS_OPTIONS,
                        index=cur_idx,
                        key=f"st_{loop_i}",
                        label_visibility="collapsed",
                    )

                    if new_s != stat and ridx is not None:
                        df.at[ridx, "Status"] = new_s
                        changed = True

        if changed:
            df.to_csv(CSV_OUTPUT, index=False)
            load_jobs_df.clear()

            try:
                from openpyxl import load_workbook

                wb = load_workbook(XLSX_OUTPUT)

                if "All Jobs" in wb.sheetnames:
                    ws = wb["All Jobs"]
                    hdrs = [c.value for c in ws[1]]

                    if "URL" in hdrs and "Status" in hdrs:
                        uc = hdrs.index("URL") + 1
                        sc_i = hdrs.index("Status") + 1
                        us = dict(zip(df["URL"], df["Status"]))

                        for r in ws.iter_rows(min_row=2, max_row=ws.max_row):
                            if r[uc - 1].value in us:
                                r[sc_i - 1].value = us[r[uc - 1].value]

                    wb.save(XLSX_OUTPUT)

            except Exception:
                pass

            st.success("Status saved to CSV and Excel.")

    except FileNotFoundError:
        st.warning("No data yet. Add roles and scrape.")

    except Exception as e:
        st.error(f"Error: {e}")