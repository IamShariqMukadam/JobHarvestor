# utils/exporter.py

import pandas as pd
import os
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter
from config import CSV_OUTPUT, XLSX_OUTPUT

# AFTER
COLUMNS = [
    "Role Searched", "Title", "Company", "Platform", "Location", "Experience Required",
    "Salary", "Date Posted", "Skills Required", "Nice To Have", 
    "URL", "Status"
]


def save_csv(jobs: list[dict]):
    os.makedirs(os.path.dirname(CSV_OUTPUT), exist_ok=True)
    df_new = pd.DataFrame(jobs, columns=COLUMNS)

    if os.path.exists(CSV_OUTPUT):
        df_existing = pd.read_csv(CSV_OUTPUT)
        df = pd.concat([df_existing, df_new], ignore_index=True)
        df.drop_duplicates(subset=["URL", "Title", "Company"], inplace=True)
    else:
        df = df_new.drop_duplicates(subset=["URL", "Title", "Company"])

    df = filter_locations(df)   # ← only line added inside save_csv

    df.to_csv(CSV_OUTPUT, index=False)
    print(f"\n  [CSV] {len(df)} jobs → {CSV_OUTPUT}")
    return df

def filter_locations(df: pd.DataFrame) -> pd.DataFrame:
    from config import TARGET_LOCATIONS
    if not TARGET_LOCATIONS:
        return df
    pattern = "|".join(TARGET_LOCATIONS)
    mask = df["Location"].str.contains(pattern, case=False, na=False)
    filtered = df[mask].copy()
    print(f"  [Filter] {len(df)} → {len(filtered)} jobs after location filter")
    return filtered

def style_header(ws, hex_color):
    fill = PatternFill("solid", fgColor=hex_color)
    font = Font(bold=True, color="FFFFFF", size=11)
    align = Alignment(horizontal="center", vertical="center")
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = align


def auto_width(ws):
    for col in ws.columns:
        max_len = max((len(str(c.value)) if c.value else 0 for c in col), default=10)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 60)


def save_xlsx(df: pd.DataFrame):
    os.makedirs(os.path.dirname(XLSX_OUTPUT), exist_ok=True)

    with pd.ExcelWriter(XLSX_OUTPUT, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="All Jobs", index=False)

        platform_summary = (
            df.groupby(["Platform", "Role Searched"])
            .size().reset_index(name="Job Count")
        )
        platform_summary.to_excel(writer, sheet_name="Platform Breakdown", index=False)

        pd.DataFrame({"Note": ["Skill frequency — added Day 2 after LLM extraction."]}) \
            .to_excel(writer, sheet_name="Skill Frequency", index=False)

        pd.DataFrame({"Note": ["Gap analysis — added Day 2 after LLM extraction."]}) \
            .to_excel(writer, sheet_name="Gap Analysis", index=False)

    wb = load_workbook(XLSX_OUTPUT)
    colors = {"All Jobs": "1F4E79", "Platform Breakdown": "375623",
              "Skill Frequency": "7B3F00", "Gap Analysis": "4A235A"}

    status_colors = {
        "To Apply": "D9E1F2", "Applied": "E2EFDA",
        "Interview": "FFF2CC", "Rejected": "FCE4D6", "Offer": "C6EFCE"
    }

    for sheet, color in colors.items():
        ws = wb[sheet]
        style_header(ws, color)
        auto_width(ws)
        ws.freeze_panes = "A2"

    ws_jobs = wb["All Jobs"]
    status_col_idx = COLUMNS.index("Status") + 1
    for row in ws_jobs.iter_rows(min_row=2, max_row=ws_jobs.max_row):
        cell = row[status_col_idx - 1]
        cell.fill = PatternFill("solid", fgColor=status_colors.get(str(cell.value), "FFFFFF"))

    wb.save(XLSX_OUTPUT)
    print(f"  [XLSX] {len(df)} jobs → {XLSX_OUTPUT}")