# main3.py — Day 3 runner
# Embeddings → Clustering → Frequency Analysis → XLSX update

import pandas as pd
import numpy as np
import json
import os
from config import CSV_OUTPUT, XLSX_OUTPUT
from clustering.embedder import get_embeddings
from clustering.clusterer import cluster, get_pca_coords
from clustering.analyzer import analyze_all


CLUSTER_REPORT = os.path.join(os.getenv("JH_SESSION_DIR","data"), "cluster_report.json")


def prepare_text(df: pd.DataFrame) -> list[str]:
    """
    Combines Skills + Role + Title into one text per job for embedding.
    More context = better clusters.
    """
    texts = []
    for _, row in df.iterrows():
        parts = [
            str(row.get("Role Searched", "")),
            str(row.get("Title", "")),
            str(row.get("Skills Required", "")),
            str(row.get("Nice To Have", ""))
        ]
        text = " ".join(p for p in parts if p not in ["N/A", "nan", ""])
        texts.append(text)
    return texts


def update_xlsx(df: pd.DataFrame, profiles: list[dict]):
    """Updates XLSX with cluster data — Sheet 3 now has per-cluster skill frequency."""
    import openpyxl
    from openpyxl.styles import PatternFill, Font, Alignment
    from openpyxl.utils import get_column_letter

    # Build cluster frequency sheet
    rows = []
    for profile in profiles:
        for skill_data in profile["top_skills"]:
            rows.append({
                "Cluster": f"{profile['cluster_id']} — {profile['label']}",
                "Skill": skill_data["skill"],
                "Count": skill_data["count"],
                "% of Cluster Jobs": f"{skill_data['percentage']}%"
            })
    cluster_skill_df = pd.DataFrame(rows)

    # Platform breakdown
    platform_df = (df.groupby(["Cluster", "Platform"])
                   .size().reset_index(name="Job Count"))

    # Gap analysis placeholder
    gap_df = pd.DataFrame({
        "Instructions": ["Paste your skills in column B — Day 4 auto-compares."],
        "Your Skills": [""]
    })

    with pd.ExcelWriter(XLSX_OUTPUT, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="All Jobs", index=False)
        platform_df.to_excel(writer, sheet_name="Platform Breakdown", index=False)
        cluster_skill_df.to_excel(writer, sheet_name="Skill Frequency", index=False)
        gap_df.to_excel(writer, sheet_name="Gap Analysis", index=False)

    # Styling
    wb = openpyxl.load_workbook(XLSX_OUTPUT)
    colors = {"All Jobs": "1F4E79", "Platform Breakdown": "375623",
              "Skill Frequency": "7B3F00", "Gap Analysis": "4A235A"}
    status_colors = {"To Apply": "D9E1F2", "Applied": "E2EFDA",
                     "Interview": "FFF2CC", "Rejected": "FCE4D6", "Offer": "C6EFCE"}

    for sheet, color in colors.items():
        ws = wb[sheet]
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor=color)
            cell.font = Font(bold=True, color="FFFFFF", size=11)
            cell.alignment = Alignment(horizontal="center")
        for col in ws.columns:
            max_len = max((len(str(c.value)) if c.value else 0 for c in col), default=10)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 60)
        ws.freeze_panes = "A2"

    ws_jobs = wb["All Jobs"]
    cols = list(df.columns)
    if "Status" in cols:
        status_idx = cols.index("Status") + 1
        for row in ws_jobs.iter_rows(min_row=2, max_row=ws_jobs.max_row):
            cell = row[status_idx - 1]
            cell.fill = PatternFill("solid", fgColor=status_colors.get(str(cell.value), "FFFFFF"))

    wb.save(XLSX_OUTPUT)
    print(f"  [XLSX] Updated → {XLSX_OUTPUT}")


def run():
    print("=" * 55)
    print("  JobHarvestor — Day 3: Embeddings + Clustering")
    print("=" * 55)

    # Load CSV
    df = pd.read_csv(CSV_OUTPUT)
    str_cols = ["Skills Required", "Nice To Have", "Experience Required", "Salary"]
    for col in str_cols:
        if col not in df.columns:
            df[col] = "N/A"
        df[col] = df[col].astype(str).replace("nan", "N/A")

    print(f"\n  Loaded {len(df)} jobs from CSV")

    # Step 1 — Embeddings
    texts = prepare_text(df)
    embeddings = get_embeddings(texts)

    # Step 2 — Clustering
    labels = cluster(embeddings)
    df["Cluster"] = labels

    # Step 3 — PCA for visualization coords
    coords = get_pca_coords(embeddings)
    df["PCA_X"] = coords[:, 0]
    df["PCA_Y"] = coords[:, 1]

    # Step 4 — Analyze clusters
    print("\n  [Analyzer] Analyzing clusters...")
    profiles = analyze_all(df)

    # Step 5 — Save cluster report
    os.makedirs("data", exist_ok=True)
    with open(CLUSTER_REPORT, "w") as f:
        json.dump(profiles, f, indent=2)
    print(f"  [Report] Saved → {CLUSTER_REPORT}")

    # Step 6 — Save updated CSV + XLSX
    df.to_csv(CSV_OUTPUT, index=False)
    print(f"  [CSV] Updated → {CSV_OUTPUT}")
    update_xlsx(df, profiles)

    # Step 7 — Print summary
    print("\n" + "=" * 55)
    print("  Day 3 Complete!")
    print("=" * 55)
    for p in profiles:
        print(f"\n  Cluster {p['cluster_id']} — {p['label']}")
        print(f"  Jobs: {p['total_jobs']}")
        print(f"  Top 5 skills:")
        for s in p["top_skills"][:5]:
            print(f"    {s['skill']:<25} {s['percentage']}%")


if __name__ == "__main__":
    run()